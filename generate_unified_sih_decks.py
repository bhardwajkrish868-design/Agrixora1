#!/usr/bin/env python3
"""
🌾 Smart India Hackathon 2026 (SIH 2026) Dual & Unified PPT Generator
Generates matching 6-slide presentations for:
1. SIH 26033: Multiple intermediaries reduce farmer earnings and increase consumer prices.
2. SIH 26091: AI-Driven Hyper-Local Business Advisory and Financial Structuring Assistant for Rural Micro-Entrepreneurs.
3. SIH Unified 26033 & 26091: Integrated Farmgate Disintermediation + Rural Enterprise Financial Structuring.
"""

import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE_PATH = r"C:\Users\krish\.gemini\antigravity\brain\a1b1e0ce-5606-49de-ad9f-d8a62ecbe1e7\.user_uploaded\media_1790529116043.pptx"
OUT_DIR = r"C:\Users\krish\.gemini\antigravity\scratch\agri-smart-platform"

# Color Palette
COLOR_NAVY = RGBColor(15, 23, 42)       # Dark Navy (#0F172A)
COLOR_EMERALD = RGBColor(16, 120, 75)   # Agri Emerald (#10784B)
COLOR_BLUE = RGBColor(26, 86, 219)      # Institutional Blue (#1A56DB)
COLOR_BODY = RGBColor(30, 41, 59)       # Slate Body (#1E293B)
COLOR_AMBER = RGBColor(180, 83, 9)      # Accent Amber (#B45309)

FONT_HEADING = "Calibri"
FONT_BODY = "Calibri"

def add_header_paragraph(tf, text, color=COLOR_EMERALD, space_before=Pt(6), font_size=Pt(12)):
    p = tf.add_paragraph()
    p.space_before = space_before
    p.space_after = Pt(2)
    run = p.add_run()
    run.text = text
    run.font.name = FONT_HEADING
    run.font.bold = True
    run.font.size = font_size
    run.font.color.rgb = color
    return p

def add_bullet_point(tf, prefix, content, color=COLOR_BODY, prefix_color=COLOR_NAVY, space_before=Pt(1.5), font_size=Pt(9.5)):
    p = tf.add_paragraph()
    p.level = 0
    p.space_before = space_before
    p.space_after = Pt(1.5)
    
    run_bullet = p.add_run()
    run_bullet.text = "• "
    run_bullet.font.name = FONT_BODY
    run_bullet.font.bold = True
    run_bullet.font.size = font_size
    run_bullet.font.color.rgb = prefix_color
    
    if prefix:
        run_pre = p.add_run()
        run_pre.text = prefix + ": "
        run_pre.font.name = FONT_BODY
        run_pre.font.bold = True
        run_pre.font.size = font_size
        run_pre.font.color.rgb = prefix_color
        
    run_content = p.add_run()
    run_content.text = content
    run_content.font.name = FONT_BODY
    run_content.font.bold = False
    run_content.font.size = font_size
    run_content.font.color.rgb = color
    return p

def format_title_shape(title_shape, title_text, font_size=Pt(19)):
    title_shape.text_frame.clear()
    p = title_shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title_text
    run.font.name = FONT_HEADING
    run.font.bold = True
    run.font.size = font_size
    run.font.color.rgb = COLOR_NAVY

def update_team_badge(slide, team_name="Ashra"):
    for shape in slide.shapes:
        if "Oval" in shape.name and shape.has_text_frame:
            shape.text_frame.clear()
            p = shape.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = team_name
            run.font.name = FONT_HEADING
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)

def build_deck(ps_id, ps_title, idea_title, filename):
    prs = pptx.Presentation(TEMPLATE_PATH)
    
    # -------------------------------------------------------------
    # SLIDE 1: TITLE PAGE
    # -------------------------------------------------------------
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text
            if "Problem Statement ID" in text or "TITLE PAGE" in text:
                tf = shape.text_frame
                tf.clear()
                
                p_title = tf.paragraphs[0]
                r_title = p_title.add_run()
                r_title.text = f"AGRIXORA: {idea_title.upper()}"
                r_title.font.bold = True
                r_title.font.size = Pt(17)
                r_title.font.color.rgb = COLOR_NAVY
                p_title.space_after = Pt(8)
                
                fields = [
                    ("Problem Statement ID", ps_id),
                    ("Problem Statement Title", ps_title),
                    ("Theme", "Agriculture, FoodTech & Rural Development"),
                    ("PS Category", "Software (AI Advisory + IoT Cold-Chain Simulation)"),
                    ("Team ID", "[Enter Registered Team ID]"),
                    ("Team Name", "Ashra"),
                    ("Working Prototype", "https://agrixora.netlify.app (or Cloudflare Edge Live)")
                ]
                
                for k, v in fields:
                    p = tf.add_paragraph()
                    p.space_after = Pt(2)
                    r1 = p.add_run()
                    r1.text = f"• {k}: "
                    r1.font.bold = True
                    r1.font.size = Pt(10)
                    r1.font.color.rgb = COLOR_NAVY
                    
                    r2 = p.add_run()
                    r2.text = str(v)
                    r2.font.bold = (k in ["Problem Statement ID", "Team Name"])
                    r2.font.size = Pt(10)
                    r2.font.color.rgb = COLOR_EMERALD if k in ["Problem Statement ID", "Working Prototype"] else COLOR_BODY

    # -------------------------------------------------------------
    # SLIDE 2: PROPOSED SOLUTION & INNOVATION
    # -------------------------------------------------------------
    s2 = prs.slides[1]
    update_team_badge(s2, "Ashra")
    for shape in s2.shapes:
        if shape == s2.shapes.title:
            format_title_shape(shape, f"IDEA TITLE — {idea_title}")
        elif shape.has_text_frame and shape != s2.shapes.title:
            tf = shape.text_frame
            tf.clear()
            
            add_header_paragraph(tf, "1. Proposed Solution & Architecture Overview", COLOR_EMERALD, Pt(2), Pt(12))
            add_bullet_point(tf, "Disintermediation Engine", "Eliminates predatory middlemen by directly linking farmers/FPOs to bulk enterprise buyers with guaranteed demand.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "AI Hyper-Local Advisory", "Assesses district-level raw material catchment to establish high-margin processing units (Dal Mill, Cold Storage, Oil Expeller).", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "4-Role Unified Platform", "Role-based portals for Farmers, Bulk Institutional Buyers, Collection Hub Operators, and Administrators.", COLOR_BODY, COLOR_NAVY)
            
            add_header_paragraph(tf, "2. Detailed Operational Workflow", COLOR_BLUE, Pt(5), Pt(12))
            add_bullet_point(tf, "Bulk Pre-Orders & Pooling", "Buyers post pre-harvest demand; smallholder farmers aggregate volumes to fulfill 500T+ contracts.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "100% Pre-Funded Escrow", "Purchase value is locked in digital multi-sig escrow prior to harvest loading; DBT payout in <2 hours.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "NABL Hub QC & Silo Storage", "Intake weighbridge verification and digital quality grading (Grade A/B) before truck dispatch.", COLOR_BODY, COLOR_NAVY)
            
            add_header_paragraph(tf, "3. Innovation & Uniqueness", COLOR_AMBER, Pt(5), Pt(12))
            add_bullet_point(tf, "Govt Subsidy Structuring", "Automated PM-FME 35% capital subsidy, AIF 3% interest relief, and 1-click Bankable DPR Dossier generation.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Bilingual Kisan Voice AI", "100% pure Hindi & English speech assistant for real-time Mandi rates and business advisory.", COLOR_BODY, COLOR_NAVY)

    # -------------------------------------------------------------
    # SLIDE 3: TECHNICAL APPROACH
    # -------------------------------------------------------------
    s3 = prs.slides[2]
    update_team_badge(s3, "Ashra")
    for shape in s3.shapes:
        if shape == s3.shapes.title:
            format_title_shape(shape, "TECHNICAL APPROACH & SYSTEM ARCHITECTURE")
        elif shape.has_text_frame and shape != s3.shapes.title:
            tf = shape.text_frame
            tf.clear()
            
            add_header_paragraph(tf, "1. Technologies & Stack Implemented", COLOR_EMERALD, Pt(2), Pt(12))
            add_bullet_point(tf, "Frontend & Client", "React 18, TypeScript, Tailwind CSS, Vite, Lucide Icons, Recharts (responsive PWA).", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Database & Cloud", "Turso Cloud Distributed LibSQL (edge-replicated <15ms response) + offline-first LocalStorage sync.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "AI & Algorithmic Engines", "Haversine Hyperlocal Radar (0-150 km), AI APMC Price Intelligence, and Cold-Chain Truck Allocation.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Security & Compliance", "Aadhaar e-KYC, Multi-Signature Escrow Vault, VAHAN Driver Compliance Registry.", COLOR_BODY, COLOR_NAVY)
            
            add_header_paragraph(tf, "2. End-to-End Implementation Flow", COLOR_BLUE, Pt(5), Pt(12))
            add_bullet_point(tf, "Step 1 [Demand / Advisory]", "Buyer creates bulk pre-order OR Farmer evaluates local processing venture via AI Advisor.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Step 2 [Pooling & Subsidy]", "Farmers pool harvest volume; Scheme engine structures 35% PM-FME subsidy & Bank DPR.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Step 3 [QC & Cold Transit]", "Intake at Collection Hub with digital grading; AI reefer dispatch with 2°C-14°C live GPS telemetry.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Step 4 [Delivery & Payout]", "Geofenced delivery verification triggers instant automated escrow DBT settlement.", COLOR_BODY, COLOR_NAVY)

    # -------------------------------------------------------------
    # SLIDE 4: FEASIBILITY AND VIABILITY
    # -------------------------------------------------------------
    s4 = prs.slides[3]
    update_team_badge(s4, "Ashra")
    for shape in s4.shapes:
        if shape == s4.shapes.title:
            format_title_shape(shape, "FEASIBILITY AND VIABILITY ANALYSIS")
        elif shape.has_text_frame and shape != s4.shapes.title:
            tf = shape.text_frame
            tf.clear()
            
            add_header_paragraph(tf, "1. Feasibility Assessment", COLOR_EMERALD, Pt(2), Pt(12))
            add_bullet_point(tf, "Technical Feasibility", "Fully working prototype with role-based workflows, Turso Cloud DB, and simulated cold telemetry already live.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Operational Feasibility", "Leverages existing physical APMC collection hubs, FCI silos, and state transporter networks (Zero Capex).", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Financial Viability", "Zero fees for farmers; nominal 1.5% platform transaction fee on corporate bulk buyers sustains operations.", COLOR_BODY, COLOR_NAVY)
            
            add_header_paragraph(tf, "2. Phased Rollout Plan", COLOR_BLUE, Pt(5), Pt(12))
            add_bullet_point(tf, "Phase 1 [Pilot]", "1 District / FPO cluster + local collection centre; validate pre-orders, fill rates, and DPR bank acceptance.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Phase 2 [Operations]", "Scale bulk demand pooling, shared transport routes, and standardized digital QC certifications.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Phase 3 [Scale & Maturity]", "Multi-district expansion, corporate buyer integrations, and full state-level logistics scaling.", COLOR_BODY, COLOR_NAVY)
            
            add_header_paragraph(tf, "3. Risk Mitigation Strategies", COLOR_AMBER, Pt(5), Pt(12))
            add_bullet_point(tf, "Low Digital Literacy", "Bilingual UI + Kisan Voice AI + on-ground FPO village coordinator onboarding support.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Payment Uncertainty", "100% upfront escrow deposit eliminates default risks; pre-dispatch NABL testing stops quality disputes.", COLOR_BODY, COLOR_NAVY)

    # -------------------------------------------------------------
    # SLIDE 5: IMPACT AND BENEFITS
    # -------------------------------------------------------------
    s5 = prs.slides[4]
    update_team_badge(s5, "Ashra")
    for shape in s5.shapes:
        if shape == s5.shapes.title:
            format_title_shape(shape, "IMPACT AND QUANTIFIABLE BENEFITS")
        elif shape.has_text_frame and shape != s5.shapes.title:
            tf = shape.text_frame
            tf.clear()
            
            add_header_paragraph(tf, "1. Stakeholder Impact Matrix", COLOR_EMERALD, Pt(2), Pt(12))
            add_bullet_point(tf, "Farmers & FPOs", "+20% to +35% higher net realization by bypassing middleman commissions; guaranteed pre-harvest demand.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Rural Micro-Entrepreneurs", "Accessible PM-FME 35% subsidies and bankable DPRs enable profitable processing units (Dal/Oil/Cold Storage).", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Bulk Institutional Buyers", "Reliable, quality-graded supply at scale with transparent pricing and real-time shipment visibility.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Consumers & Society", "10% to 18% lower retail food prices; 60% reduction in post-harvest perishable wastage.", COLOR_BODY, COLOR_NAVY)
            
            add_header_paragraph(tf, "2. Triple-Bottom-Line Benefits", COLOR_BLUE, Pt(5), Pt(12))
            add_bullet_point(tf, "Economic", "Higher farmgate income, zero payment defaults, creation of local rural agro-processing jobs.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Social", "Inclusion of smallholder & women farmers; trust built through transparent digital QC and instant DBT payouts.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Environmental", "Cold-chain integration prevents perishable rotting; optimized transport routes reduce carbon footprint.", COLOR_BODY, COLOR_NAVY)

    # -------------------------------------------------------------
    # SLIDE 6: RESEARCH AND REFERENCES
    # -------------------------------------------------------------
    s6 = prs.slides[5]
    update_team_badge(s6, "Ashra")
    for shape in s6.shapes:
        if shape == s6.shapes.title:
            format_title_shape(shape, "RESEARCH GROUNDING AND REFERENCES")
        elif shape.has_text_frame and shape != s6.shapes.title:
            tf = shape.text_frame
            tf.clear()
            
            add_header_paragraph(tf, "1. Live Prototype Implementation Evidence", COLOR_EMERALD, Pt(2), Pt(12))
            add_bullet_point(tf, "Live Working System", "Agrixora Platform (React 18 / TypeScript / Vite / Turso Distributed LibSQL Cloud DB).", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Verified Features", "Bulk Pre-Orders, AI Business & Subsidy Advisory, Digital QC Labs, GPS Reefer Fleet, Kisan Voice AI.", COLOR_BODY, COLOR_NAVY)
            
            add_header_paragraph(tf, "2. Government Policies & Institutional Standards", COLOR_BLUE, Pt(5), Pt(12))
            add_bullet_point(tf, "MoA&FW Guidelines", "Committee on Doubling Farmers' Income (Post-production & direct marketing framework).", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "MoFPI PM-FME Scheme", "Pradhan Mantri Formalisation of Micro Food Processing Enterprises (35% Credit-Linked Subsidy).", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "Agriculture Infra Fund (AIF)", "Central Sector Scheme for 3% interest subvention & credit guarantee coverage up to ₹2.00 Cr.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "National Agriculture Market", "e-NAM & AGMARKNET integration for daily transparent spot price benchmarking.", COLOR_BODY, COLOR_NAVY)
            
            add_header_paragraph(tf, "3. Academic & Technical Research Citations", COLOR_AMBER, Pt(5), Pt(12))
            add_bullet_point(tf, "ICAR-CIPHET Study", "Comprehensive assessment of harvest and post-harvest losses in major Indian agricultural crops.", COLOR_BODY, COLOR_NAVY)
            add_bullet_point(tf, "NITI Aayog Policy", "Report on Agritech Innovations, Farmer-Buyer Linkages, and Decentralized Cold Supply Chains.", COLOR_BODY, COLOR_NAVY)

    out_path = os.path.join(OUT_DIR, filename)
    prs.save(out_path)
    print(f"Generated presentation: {out_path}")
    return out_path

if __name__ == "__main__":
    # 1. Deck for SIH 26033
    build_deck(
        ps_id="26033",
        ps_title="Multiple intermediaries reduce farmer earnings and increase consumer prices.",
        idea_title="FARM2FUTURE (Agrixora: Direct Marketplace & Pre-Order Supply Chain)",
        filename="SIH_26033_Submission.pptx"
    )

    # 2. Deck for SIH 26091
    build_deck(
        ps_id="26091",
        ps_title="AI-Driven Hyper-Local Business Advisory and Financial Structuring Assistant for Rural Micro-Entrepreneurs",
        idea_title="FARM2FUTURE (Agrixora: Rural Enterprise Advisory & Subsidy Structuring)",
        filename="SIH_26091_Submission.pptx"
    )

    # 3. Unified Deck for Both
    build_deck(
        ps_id="26033 & 26091 (Dual Unified Solution)",
        ps_title="Direct Farmgate Pre-Orders + AI Rural Micro-Enterprise & Subsidy Structuring",
        idea_title="FARM2FUTURE (Agrixora: Complete Agri Supply & Enterprise Ecosystem)",
        filename="SIH_Unified_26033_26091_Submission.pptx"
    )
